"""Phase 3 tests: CSVDataSource and the dataset-agnostic validator helpers."""

from __future__ import annotations

from io import BytesIO

import openpyxl
import pandas as pd
import pytest

from src.ingestion.csv_loader import CSVDataSource
from src.ingestion.mapper import INTERNAL_FIELDS
from src.ingestion.template import build_template_workbook
from src.ingestion.validator import (
    compute_missingness,
    count_duplicate_rows,
    detect_datetime_columns,
    detect_id_like_columns,
    validate_required_columns,
)


@pytest.fixture()
def sample_csv(tmp_path):
    path = tmp_path / "sample_assets.csv"
    df = pd.DataFrame(
        {
            "Asset ID": ["4212-CR-001", "4212-PU-002", "4213-CV-202"],
            "Category": ["Crusher", "Pump", None],
            "Overdue PMs": [2, 0, 5],
            "Last Completed PM": ["2025-01-03", "2025-01-10", None],
        }
    )
    df.to_csv(path, index=False)
    return path


def test_csv_inspect_basic_profile(sample_csv):
    source = CSVDataSource(sample_csv)
    result = source.inspect()
    assert result.row_count == 3
    assert result.column_count == 4
    assert result.sheet_names is None


def test_csv_validate_missing_file(tmp_path):
    source = CSVDataSource(tmp_path / "missing.csv")
    problems = source.validate()
    assert any("not found" in p for p in problems)


def test_csv_validate_wrong_extension(tmp_path):
    bad = tmp_path / "data.xlsx"
    bad.write_text("not a csv")
    source = CSVDataSource(bad)
    problems = source.validate()
    assert any("Unsupported file extension" in p for p in problems)


def test_csv_load_roundtrip(sample_csv):
    source = CSVDataSource(sample_csv)
    df = source.load()
    assert len(df) == 3
    assert "Asset ID" in df.columns


def test_detect_datetime_columns_by_name_and_content():
    df = pd.DataFrame(
        {
            "Last Completed PM": ["2025-01-03", "2025-01-10", "2025-02-01"],
            "Asset ID": ["A1", "A2", "A3"],
            "Overdue PMs": [1, 2, 3],
        }
    )
    found = detect_datetime_columns(df)
    assert "Last Completed PM" in found
    assert "Overdue PMs" not in found


def test_detect_id_like_columns_by_name():
    df = pd.DataFrame({"Asset ID": ["A1", "A2", "A3"], "Downtime": [1.0, 2.0, 3.0]})
    found = detect_id_like_columns(df)
    assert "Asset ID" in found


def test_compute_missingness():
    df = pd.DataFrame({"a": [1, None, 3, None], "b": [1, 2, 3, 4]})
    missing = compute_missingness(df)
    assert missing["a"] == 50.0
    assert missing["b"] == 0.0


def test_count_duplicate_rows():
    df = pd.DataFrame({"a": [1, 1, 2], "b": [1, 1, 2]})
    assert count_duplicate_rows(df) == 1


def test_validate_required_columns_reports_missing():
    df = pd.DataFrame({"a": [1]})
    problems = validate_required_columns(df, ["a", "b"])
    assert len(problems) == 1
    assert "b" in problems[0]


def test_build_template_workbook_is_a_valid_readable_xlsx():
    workbook_bytes = build_template_workbook()
    wb = openpyxl.load_workbook(BytesIO(workbook_bytes))
    assert wb.sheetnames == ["Data", "Field Reference"]


def test_template_data_sheet_header_matches_internal_fields():
    workbook_bytes = build_template_workbook()
    wb = openpyxl.load_workbook(BytesIO(workbook_bytes))
    ws = wb["Data"]
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header == list(INTERNAL_FIELDS.keys())


def test_template_field_reference_sheet_flags_required_fields():
    workbook_bytes = build_template_workbook()
    wb = openpyxl.load_workbook(BytesIO(workbook_bytes))
    ws = wb["Field Reference"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    required_field_rows = [r for r in rows if r[0] == "machine_id"]
    assert len(required_field_rows) == 1
    assert required_field_rows[0][1] == "Yes"


def test_template_is_uploadable_through_the_real_excel_loader(tmp_path):
    # The template must round-trip through the actual ingestion path, not
    # just be well-formed in isolation.
    from src.ingestion.excel_loader import ExcelDataSource

    path = tmp_path / "template.xlsx"
    path.write_bytes(build_template_workbook())

    source = ExcelDataSource(path, sheet_name="Data")
    assert source.validate() == []
    df = source.load()
    assert list(df.columns) == list(INTERNAL_FIELDS.keys())
